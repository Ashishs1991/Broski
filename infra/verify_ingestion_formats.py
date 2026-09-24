"""Run inside worker image: python /checks/verify_ingestion_formats.py."""
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
from docx import Document

from app.extract import extract


with TemporaryDirectory() as directory:
    root = Path(directory)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Policy limit is 500000 rupees.")
    pdf.save(root / "digital.pdf")
    digital = extract(root / "digital.pdf", 10)
    assert digital[0]["page_number"] == 1
    assert "500000" in digital[0]["text"]

    image = page.get_pixmap(matrix=pymupdf.Matrix(3, 3)).tobytes("png")
    (root / "scan.png").write_bytes(image)
    assert "500000" in " ".join(c["text"] for c in extract(root / "scan.png", 10))

    scanned = pymupdf.open()
    scan_page = scanned.new_page()
    scan_page.insert_image(scan_page.rect, stream=image)
    scanned.save(root / "scanned.pdf")
    assert "500000" in " ".join(c["text"] for c in extract(root / "scanned.pdf", 10))

    word = Document()
    word.add_heading("Coverage", level=1)
    word.add_paragraph("The policy has a deductible.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Type"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Health"
    table.cell(1, 1).text = "500000"
    word.save(root / "policy.docx")
    chunks = extract(root / "policy.docx", 10)
    assert any("Coverage" == c["section_path"] for c in chunks)
    assert any("| Type | Limit |" in c["text"] for c in chunks)

    (root / "note.md").write_text("# Note\n\nImportant observation.", encoding="utf-8")
    assert extract(root / "note.md", 10)[0]["section_path"] == "Note"

    (root / "bad.pdf").write_bytes(b"%PDF-corrupt")
    try:
        extract(root / "bad.pdf", 10)
    except Exception:
        pass
    else:
        raise AssertionError("Corrupt PDF was accepted")

print("PASS: digital PDF, scanned PDF, image, DOCX heading/table, note, corrupt PDF")
